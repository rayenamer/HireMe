#!/usr/bin/env python3
"""Extract plain text from a .docx file.

Usage: python3 docx_to_text.py <path-to-docx>

Prints extracted text to stdout, paragraph order preserved, tables
flattened to tab-separated rows. Tries python-docx first; falls back
to a stdlib-only OOXML reader if python-docx isn't installed, so this
script has zero hard dependencies.
"""
import sys
import zipfile
import xml.etree.ElementTree as ET

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def extract_with_python_docx(path):
    import docx  # type: ignore

    doc = docx.Document(path)
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        is_list = para.style is not None and "List" in (para.style.name or "")
        lines.append(f"- {text}" if is_list else text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append("\t".join(cells))
    return "\n".join(lines)


def extract_with_stdlib(path):
    with zipfile.ZipFile(path) as z:
        xml_bytes = z.read("word/document.xml")
    root = ET.fromstring(xml_bytes)
    body = root.find(f"{W_NS}body")
    if body is None:
        return ""

    def paragraph_text(p_elem):
        texts = [t.text or "" for t in p_elem.iter(f"{W_NS}t")]
        has_num_pr = p_elem.find(f".//{W_NS}numPr") is not None
        style_elem = p_elem.find(f"{W_NS}pPr/{W_NS}pStyle")
        style_val = style_elem.get(f"{W_NS}val", "") if style_elem is not None else ""
        is_list = has_num_pr or "list" in style_val.lower()
        joined = "".join(texts).strip()
        if not joined:
            return None
        return f"- {joined}" if is_list else joined

    lines = []
    for elem in body:
        tag = elem.tag
        if tag == f"{W_NS}p":
            line = paragraph_text(elem)
            if line:
                lines.append(line)
        elif tag == f"{W_NS}tbl":
            for row in elem.iter(f"{W_NS}tr"):
                cells = []
                for cell in row.iter(f"{W_NS}tc"):
                    cell_texts = [t.text or "" for t in cell.iter(f"{W_NS}t")]
                    cells.append("".join(cell_texts).strip())
                if any(cells):
                    lines.append("\t".join(cells))
    return "\n".join(lines)


def main():
    if len(sys.argv) != 2:
        print("Usage: python3 docx_to_text.py <path-to-docx>", file=sys.stderr)
        sys.exit(1)
    path = sys.argv[1]
    try:
        text = extract_with_python_docx(path)
    except ImportError:
        text = extract_with_stdlib(path)
    print(text)


if __name__ == "__main__":
    main()
