#!/usr/bin/env python3
"""Generate an ATS-friendly .docx from the structured Markdown convention
documented in SKILL.md:

    # Full Name
    contact line
    <blank>
    ## Section Heading
    ### Job Title — Company (Mon YYYY - Mon YYYY)
    - bullet
    plain paragraph

Usage: python3 text_to_docx.py <input.md> <output.docx>

Requires python-docx (pip install python-docx). Exits with a clear,
non-zero-status message if it isn't installed, rather than crashing
with a traceback - the calling skill treats that as a soft failure
and falls back to the Markdown output alone.
"""
import sys


def main():
    if len(sys.argv) != 3:
        print("Usage: python3 text_to_docx.py <input.md> <output.docx>", file=sys.stderr)
        sys.exit(1)

    src, dest = sys.argv[1], sys.argv[2]

    try:
        import docx
        from docx.shared import Pt
    except ImportError:
        print(
            "python-docx is not installed. Run: pip install python-docx\n"
            "The Markdown version of the CV was still generated and is ATS-safe on its own.",
            file=sys.stderr,
        )
        sys.exit(2)

    with open(src, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()

    document = docx.Document()

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for raw_line in lines:
        line = raw_line.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
        elif line.startswith("# "):
            heading = document.add_heading(line[2:].strip(), level=0)
        elif line.lstrip().startswith("- ") or line.lstrip().startswith("* "):
            text = line.lstrip()[2:].strip()
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(line.strip())

    document.save(dest)
    print(f"Wrote {dest}")


if __name__ == "__main__":
    main()
